import csv
import hashlib
import json
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Consignment, ConsignmentIssue, ConsignmentLine, Marketplace
from app.services.amazon_consignment_service import AmazonConsignmentService
from app.services.consignment_matching import ConsignmentMatcher, normalized
from app.services.consignment_validation import ConsignmentValidationService
from app.services.flipkart_consignment_service import FlipkartQuantityMatcher

AMAZON_HEADERS = {"merchant sku", "seller sku", "sku", "asin", "fnsku", "shipped"}
FLIPKART_HEADERS = {"fsn", "sku", "sku id", "quantity sent", "quantity", "qty sent", "qty"}


@dataclass(frozen=True)
class ParsedConsignment:
    filename: str
    sheet: str
    detected_type: str
    header_row: int
    recognized_columns: list[str]
    rows: list[tuple[int, dict[str, Any]]]
    warnings: list[str]
    sha256: str


def _matrices(filename: str, payload: bytes):
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        try: text = payload.decode("utf-8-sig")
        except UnicodeDecodeError: text = payload.decode("cp1252")
        yield "CSV", [list(row) for row in csv.reader(StringIO(text))]
    elif suffix in {".xlsx", ".xlsm"}:
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_vba=False)
        try:
            for name in book.sheetnames: yield name, [list(row) for row in book[name].iter_rows(values_only=True)]
        finally: book.close()
    else: raise ValueError("Only CSV, XLSX, and XLSM files are supported")


def parse_consignment(filename: str, payload: bytes, marketplace: Marketplace) -> ParsedConsignment:
    expected = AMAZON_HEADERS if marketplace == Marketplace.AMAZON else FLIPKART_HEADERS
    candidates = []
    for sheet, matrix in _matrices(filename, payload):
        for index, values in enumerate(matrix[:20]):
            headers = [" ".join(str(v or "").strip().split()).casefold() for v in values]
            recognized = sorted({v for v in headers if v in expected})
            has_quantity = any(v in {"shipped", "quantity sent", "quantity", "qty sent", "qty"} for v in headers)
            has_identity = any(v in {"merchant sku", "seller sku", "sku", "sku id", "asin", "fnsku", "fsn"} for v in headers)
            if has_quantity and has_identity:
                nonempty = sum(any(v not in (None, "") for v in row) for row in matrix[index + 1:])
                candidates.append((len(recognized) * 100 + nonempty, sheet, index, values, matrix, recognized))
    if not candidates: raise ValueError("Could not detect a supported consignment quantity sheet")
    _, sheet, header_index, headers, matrix, recognized = max(candidates, key=lambda row: row[0])
    rows = []
    for number, values in enumerate(matrix[header_index + 1:], start=header_index + 2):
        if any(value not in (None, "") for value in values):
            rows.append((number, {str(headers[i] or f"column_{i + 1}"): value for i, value in enumerate(values) if i < len(headers)}))
    kind = "Amazon Consignment" if marketplace == Marketplace.AMAZON else "Flipkart Quantity"
    return ParsedConsignment(filename, sheet, kind, header_index + 1, recognized, rows, [], hashlib.sha256(payload).hexdigest())


class ConsignmentImportService:
    def __init__(self, db: Session): self.db = db

    def apply(self, consignment: Consignment, parsed: ParsedConsignment) -> dict[str, int]:
        matcher = ConsignmentMatcher(self.db, consignment.account_id)
        normalizer = AmazonConsignmentService(matcher) if consignment.marketplace == Marketplace.AMAZON else FlipkartQuantityMatcher(matcher)
        counts = {"matched": 0, "blocked": 0, "unmatched": 0, "ambiguous": 0, "total_print_quantity": 0}
        seen: dict[str, ConsignmentLine] = {}
        consignment.source_file_name, consignment.source_file_sha256 = parsed.filename, parsed.sha256
        consignment.source_type, consignment.status = parsed.detected_type, "open"
        for source_row, raw in parsed.rows:
            key = hashlib.sha256(json.dumps({str(k).casefold(): normalized(v) for k, v in raw.items()}, sort_keys=True).encode()).hexdigest()
            try:
                values = normalizer.normalize(raw)
            except ValueError as exc:
                values = {"print_quantity": 0, "source_quantity": None, "product": None, "match_status": "unmatched", "match_method": None, "match_error": str(exc)}
            product = values.pop("product")
            match_error = values.pop("match_error")
            line = ConsignmentLine(consignment=consignment, source_row=source_row, source_line_key=key,
                selected_for_print=False, net_quantity_value=1, net_quantity_unit="N", raw_row=raw, **values)
            if product:
                line.product = product
                line.title_snapshot = line.title_snapshot or product.title
                line.brand_snapshot, line.category_snapshot, line.format_key = product.brand, product.category, product.category
                line.mrp_catalog, line.mrp_source = product.mrp, "catalog" if product.mrp is not None else "missing"
                identifiers = {item.kind.casefold(): item.value for item in product.identifiers}
                line.asin, line.fnsku, line.fsn, line.listing_id = line.asin or identifiers.get("asin"), line.fnsku or identifiers.get("fnsku"), line.fsn or identifiers.get("fsn"), identifiers.get("listing_id")
            self.db.add(line); self.db.flush()
            ConsignmentValidationService(self.db).validate(line)
            if match_error and match_error != "MISSING_CATALOG_MATCH":
                self.db.add(ConsignmentIssue(consignment_id=consignment.id, consignment_line_id=line.id, severity="blocking", code=match_error, message="Catalog matching or quantity validation failed."))
            if key in seen:
                for duplicate in (seen[key], line):
                    self.db.add(ConsignmentIssue(consignment_id=consignment.id, consignment_line_id=duplicate.id, severity="blocking", code="DUPLICATE_CONSIGNMENT_LINE", message="An identical row occurs more than once; quantities were not combined."))
                    duplicate.workflow_state = "blocked"
            else: seen[key] = line
            line.error_count = len([issue for issue in self.db.new if isinstance(issue, ConsignmentIssue) and issue.consignment_line_id == line.id])
            if line.workflow_state == "blocked": counts["blocked"] += 1
            counts[line.match_status] += 1
            counts["total_print_quantity"] += max(line.print_quantity, 0)
        self.db.flush()
        return counts
