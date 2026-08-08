import csv
import hashlib
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.normalization import NormalizedProduct, detect_header_row, normalize_row


@dataclass(slots=True)
class ParsedCatalog:
    detected_type: str
    sheet_name: str
    header_row: int
    headers: list[Any]
    rows: list[tuple[int, dict, NormalizedProduct]]
    file_sha256: str
    file_size: int
    warnings: list[str] = field(default_factory=list)


def file_fingerprint(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _csv_matrix(payload: bytes) -> list[list[Any]]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = payload.decode("cp1252")
    return [list(row) for row in csv.reader(StringIO(text))]


def _sheet_score(matrix: list[list[Any]], marketplace: str) -> tuple[int, int] | None:
    try:
        header_index = detect_header_row(matrix, marketplace)
    except ValueError:
        return None
    nonempty = sum(1 for row in matrix[header_index + 1 :] if any(value not in (None, "") for value in row))
    technical_bonus = sum("#1." in str(value or "") for value in matrix[header_index]) * 20
    return technical_bonus + min(nonempty, 1000), header_index


def _best_workbook_sheet(payload: bytes, marketplace: str):
    cached_book = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_vba=False)
    formula_book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_vba=False)
    candidates = []
    for sheet_name in cached_book.sheetnames:
        cached_matrix = [list(row) for row in cached_book[sheet_name].iter_rows(values_only=True)]
        score = _sheet_score(cached_matrix, marketplace)
        if score:
            formula_matrix = [list(row) for row in formula_book[sheet_name].iter_rows(values_only=True)]
            candidates.append((score[0], sheet_name, score[1], cached_matrix, formula_matrix))
    cached_book.close()
    formula_book.close()
    if not candidates:
        raise ValueError("Could not detect a supported worksheet schema")
    return max(candidates, key=lambda item: item[0])


def _formula_warnings(cached: list[list[Any]], formulas: list[list[Any]], header_index: int) -> list[str]:
    warnings = []
    required_aliases = {"sku", "seller sku", "merchant sku", "sku id", "mrp", "maximum retail price"}
    headers = [str(value or "").strip().casefold() for value in cached[header_index]]
    required_columns = {index for index, value in enumerate(headers) if value in required_aliases or value == "contribution_sku#1.value" or "maximum_retail_price" in value}
    for row_index in range(header_index + 1, min(len(cached), len(formulas))):
        for column_index in required_columns:
            cached_value = cached[row_index][column_index] if column_index < len(cached[row_index]) else None
            formula_value = formulas[row_index][column_index] if column_index < len(formulas[row_index]) else None
            if cached_value is None and isinstance(formula_value, str) and formula_value.startswith("="):
                warnings.append(f"FORMULA_RESULT_MISSING row {row_index + 1}, column {column_index + 1}")
    return warnings


def parse_catalog(filename: str, payload: bytes, marketplace: str, source_category: str | None = None) -> ParsedCatalog:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        matrix = _csv_matrix(payload)
        header_index = detect_header_row(matrix, marketplace)
        sheet_name, formula_warnings = "CSV", []
    elif suffix in {".xlsx", ".xlsm"}:
        _, sheet_name, header_index, matrix, formula_matrix = _best_workbook_sheet(payload, marketplace)
        formula_warnings = _formula_warnings(matrix, formula_matrix, header_index)
    else:
        raise ValueError("Only CSV, XLSX, and XLSM files are supported")

    headers = matrix[header_index]
    parsed = []
    template = Path(filename).stem
    formula_rows = {int(warning.split()[2].rstrip(",")) for warning in formula_warnings}
    for source_index, values in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if source_index in formula_rows:
            continue
        raw = {str(headers[index] or f"column_{index + 1}"): value for index, value in enumerate(values) if index < len(headers)}
        parsed.append((source_index, raw, normalize_row(headers, values, marketplace, source_template=template, source_category=source_category)))
    detected = "Amazon Listing Template" if marketplace == "amazon" else "Flipkart QC Catalog"
    return ParsedCatalog(detected, sheet_name, header_index + 1, headers, parsed, file_fingerprint(payload), len(payload), formula_warnings)
